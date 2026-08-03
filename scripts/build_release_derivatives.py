"""Cognitive Industries — Les Industries Cognitives
Project: AI-IDP / AegisTrace
Author and Intellectual Property Owner: Pierre-Edward Procyk
Copyright: © 2026 Pierre-Edward Procyk. All rights reserved.
File: scripts/build_release_derivatives.py
Purpose: Regenerate current public PDF and DOCX derivatives from canonical Markdown sources.
Classification: build tooling
Version: 2.0.0
Last Material Revision: 2026-08-01
Licence Status: No licence selected unless approved in writing by Pierre-Edward Procyk.
"""

from __future__ import annotations

import argparse
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_government_pdfs import (
    BRAND_LOGO,
    PROJECT_ROOT,
    DocumentSpec,
    _parse_markdown,
    _register_fonts,
    build_document,
)


NAVY = "0F1728"
GOLD = "B89A5E"
IVORY = "FBF7EE"
BODY = "182033"
RELEASE_VERSION = "2.0.0"
RELEASE_DATE = "2026-08-01"

FULL_CONTACT_EN = (
    "Pierre-Edward Procyk",
    "Founder / CEO",
    "Cognitive Industries — Les Industries Cognitives",
    "Saguenay, Québec, Canada",
    "p.procyk.media@gmail.com",
    "p.1o9.cognitive@outlook.com",
    "",
    "LinkedIn: linkedin.com/in/pierre-edward-procyk-223b75305",
)
BRIEF_CONTACT_EN = FULL_CONTACT_EN[:5]
BRIEF_CONTACT_FR = (
    "Pierre-Edward Procyk",
    "Fondateur / PDG",
    "Cognitive Industries — Les Industries Cognitives",
    "Saguenay, Québec, Canada",
    "p.procyk.media@gmail.com",
)


def _required(*extra: str) -> tuple[str, ...]:
    return (RELEASE_VERSION, RELEASE_DATE, *extra)


def _forbidden(*extra: str) -> tuple[str, ...]:
    return (
        "1.3.0",
        "1.1.0",
        "1.0.0",
        "2026-07-28",
        "2026-07-24",
        *extra,
    )


PDF_DOCUMENTS = {
    "proposal": DocumentSpec(
        key="proposal",
        source=PROJECT_ROOT / "government" / "CANADIAN_NATIONAL_PROJECT_PROPOSAL.md",
        output=PROJECT_ROOT / "government" / "CANADIAN_NATIONAL_PROJECT_PROPOSAL.pdf",
        title="Canadian National Project Proposal",
        subtitle=(
            "AI-IDP — Universal AI Identity, Delegation, Provenance, Traceability, "
            "Quality, Accountability, and Permanent Audit Standard"
        ),
        author_role="Founder / CEO",
        document_id="AI-IDP-GOV-PROP-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=FULL_CONTACT_EN,
        cover_heading_count=2,
        required_text=_required("113"),
        forbidden_text=_forbidden("89 passing tests"),
        minimum_pages=4,
    ),
    "white-paper": DocumentSpec(
        key="white-paper",
        source=PROJECT_ROOT / "government" / "CANADIAN_POLICY_WHITE_PAPER.md",
        output=PROJECT_ROOT / "government" / "CANADIAN_POLICY_WHITE_PAPER.pdf",
        title="Canadian Policy White Paper — AI-IDP",
        subtitle=(
            "A Universal Framework for Persistent AI Actor Identity, Permanent Traceability, "
            "Delegation, Quality Assurance, and Accountable AI Operation in Canada"
        ),
        author_role="Founder / CEO",
        document_id="AI-IDP-WP-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        required_text=_required("113"),
        forbidden_text=_forbidden("89 passing tests"),
        minimum_pages=3,
    ),
    "charter": DocumentSpec(
        key="charter",
        source=PROJECT_ROOT / "government" / "CHARTER_ANALYSIS.md",
        output=PROJECT_ROOT / "government" / "CHARTER_ANALYSIS.pdf",
        title="Charter Analysis",
        subtitle="Canadian Charter of Rights and Freedoms implications of AI-IDP",
        author_role="Founder / CEO",
        document_id="AI-IDP-CHR-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=2,
    ),
    "jurisdiction": DocumentSpec(
        key="jurisdiction",
        source=PROJECT_ROOT / "government" / "FEDERAL_PROVINCIAL_JURISDICTION_ANALYSIS.md",
        output=PROJECT_ROOT / "government" / "FEDERAL_PROVINCIAL_JURISDICTION_ANALYSIS.pdf",
        title="Federal-Provincial Jurisdiction Analysis",
        subtitle="Constitutional jurisdiction over AI agent identity and traceability in Canada",
        author_role="Founder / CEO",
        document_id="AI-IDP-JUR-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=2,
    ),
    "privacy": DocumentSpec(
        key="privacy",
        source=PROJECT_ROOT / "government" / "PRIVACY_AND_HUMAN_RIGHTS_ANALYSIS.md",
        output=PROJECT_ROOT / "government" / "PRIVACY_AND_HUMAN_RIGHTS_ANALYSIS.pdf",
        title="Privacy and Human Rights Analysis",
        subtitle="Privacy and human-rights implications of AI-IDP",
        author_role="Founder / CEO",
        document_id="AI-IDP-PHR-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=2,
    ),
    "impact-fr": DocumentSpec(
        key="impact-fr",
        source=PROJECT_ROOT / "impact" / "IMPACT_CANADIEN_FR.md",
        output=PROJECT_ROOT / "impact" / "IMPACT_CANADIEN_FR.pdf",
        title="Profil d'impact canadien",
        subtitle="Analyse d'impact (affaires, RH, société) — résumé français",
        author_role="Fondateur / PDG",
        document_id="AI-IDP-IMPFR-2.0.0",
        classification="Public",
        language="Français",
        rights_text="Tous droits réservés.",
        contact_lines=BRIEF_CONTACT_FR,
        minimum_pages=2,
    ),
    "business": DocumentSpec(
        key="business",
        source=PROJECT_ROOT / "impact" / "business-and-operations" / "BUSINESS_AND_OPERATIONS_IMPACT_REPORT.md",
        output=PROJECT_ROOT / "impact" / "business-and-operations" / "BUSINESS_AND_OPERATIONS_IMPACT_REPORT.pdf",
        title="Business and Operations Impact Report",
        subtitle="AI-IDP / AegisTrace — Adoption Impact Analysis for Canadian Business and Operations",
        author_role="Founder / CEO",
        document_id="AI-IDP-BIZ-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=4,
    ),
    "hr": DocumentSpec(
        key="hr",
        source=PROJECT_ROOT / "impact" / "human-resources" / "HR_AND_LABOUR_IMPACT_REPORT.md",
        output=PROJECT_ROOT / "impact" / "human-resources" / "HR_AND_LABOUR_IMPACT_REPORT.pdf",
        title="HR and Labour Impact Report",
        subtitle="AI-IDP / AegisTrace — Human Resources and Labour Implications",
        author_role="Founder / CEO",
        document_id="AI-IDP-HR-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=3,
    ),
    "societal": DocumentSpec(
        key="societal",
        source=PROJECT_ROOT / "impact" / "societal" / "SOCIETAL_IMPACT_ASSESSMENT.md",
        output=PROJECT_ROOT / "impact" / "societal" / "SOCIETAL_IMPACT_ASSESSMENT.pdf",
        title="Societal Impact Assessment",
        subtitle="AI-IDP / AegisTrace — Societal Implications for Canada",
        author_role="Founder / CEO",
        document_id="AI-IDP-SOC-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=3,
    ),
    "technical": DocumentSpec(
        key="technical",
        source=PROJECT_ROOT / "technical" / "AEGISTRACE_TECHNICAL_ARCHITECTURE.md",
        output=PROJECT_ROOT / "technical" / "AEGISTRACE_TECHNICAL_ARCHITECTURE.pdf",
        title="AegisTrace Technical Architecture",
        subtitle="Reference Implementation of the AI-IDP Standard",
        author_role="Founder / CEO",
        document_id="AEGISTRACE-ARCH-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=BRIEF_CONTACT_EN,
        minimum_pages=5,
    ),
    "technical-fr": DocumentSpec(
        key="technical-fr",
        source=PROJECT_ROOT / "technical" / "ARCHITECTURE_TECHNIQUE_FR.md",
        output=PROJECT_ROOT / "technical" / "ARCHITECTURE_TECHNIQUE_FR.pdf",
        title="Architecture technique d'AegisTrace",
        subtitle="Implémentation de référence de la norme AI-IDP",
        author_role="Fondateur / PDG",
        document_id="AI-IDP-ARCHFR-2.0.0",
        classification="Public",
        language="Français",
        rights_text="Tous droits réservés.",
        contact_lines=BRIEF_CONTACT_FR,
        minimum_pages=3,
    ),
    "university": DocumentSpec(
        key="university",
        source=PROJECT_ROOT / "university" / "UNIVERSITY_RESEARCH_REPORT.md",
        output=PROJECT_ROOT / "university" / "UNIVERSITY_RESEARCH_REPORT.pdf",
        title="Identity Before Autonomy",
        subtitle=(
            "A Universal Framework for Persistent AI Actor Identity, Permanent Traceability, "
            "Delegation, Quality Assurance, and Accountable AI Operation in Canada"
        ),
        author_role="Author",
        document_id="AI-IDP-UNIV-2.0.0",
        classification="Public",
        language="English",
        rights_text="All rights reserved.",
        contact_lines=(
            "Pierre-Edward Procyk",
            "Cognitive Industries — Les Industries Cognitives",
            "Saguenay, Québec, Canada",
            "Correspondence: p.procyk.media@gmail.com",
        ),
        required_text=_required("113", "25 specification documents", "14 JSON schemas"),
        forbidden_text=_forbidden("89 passing tests"),
        minimum_pages=6,
    ),
}


@dataclass(frozen=True)
class DocxSpec:
    """Define one canonical Markdown-to-DOCX delivery contract."""

    key: str
    reference: Path
    pdf_spec: DocumentSpec
    cover_page_break_index: int
    cover_updates: tuple[tuple[int, str], ...]


DOCX_DOCUMENTS = {
    "proposal-docx": DocxSpec(
        key="proposal-docx",
        reference=PROJECT_ROOT / "government" / "CANADIAN_NATIONAL_PROJECT_PROPOSAL.docx",
        pdf_spec=PDF_DOCUMENTS["proposal"],
        cover_page_break_index=18,
        cover_updates=(
            (0, "Canadian National Project Proposal"),
            (1, PDF_DOCUMENTS["proposal"].subtitle),
            (12, "Document ID: AI-IDP-GOV-PROP-2.0.0"),
            (13, "Version: 2.0.0"),
            (14, "Date: 2026-08-01"),
            (15, "Classification: Public"),
            (16, "Status: Submission-ready. Not submitted."),
        ),
    ),
    "university-docx": DocxSpec(
        key="university-docx",
        reference=PROJECT_ROOT / "university" / "UNIVERSITY_RESEARCH_REPORT.docx",
        pdf_spec=PDF_DOCUMENTS["university"],
        cover_page_break_index=15,
        cover_updates=(
            (0, "Identity Before Autonomy"),
            (1, PDF_DOCUMENTS["university"].subtitle),
            (10, "Document ID: AI-IDP-UNIV-2.0.0"),
            (11, "Version: 2.0.0"),
            (12, "Date: 2026-08-01"),
            (13, "Classification: Public"),
        ),
    ),
}


def _set_run_font(run, name: str, size: float, colour: str, bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(colour)
    run.bold = bold


def _replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)


def _append_inline(paragraph, text: str) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*|\[[^\]]+\]\([^)]+\))")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            paragraph.add_run(f"{link.group(1)} ({link.group(2)})" if link else token)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_table_geometry(table, widths_twips: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths_twips)))
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for cell_width in widths_twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(cell_width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = int(widths_twips[index] * 635)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(widths_twips[index]))


def _add_table(document: DocumentObject, rows: Iterable[tuple[str, ...]]) -> None:
    materialized = tuple(rows)
    column_count = max(len(row) for row in materialized)
    padded = tuple(row + ("",) * (column_count - len(row)) for row in materialized)
    table = document.add_table(rows=len(padded), cols=column_count)
    table.style = "Table Grid"
    section = document.sections[-1]
    total_twips = int((section.page_width - section.left_margin - section.right_margin) / 635)
    weights = [max(8, max(len(re.sub(r"[*_`]", "", row[index])) for row in padded)) for index in range(column_count)]
    weight_total = sum(weights)
    widths = [max(720, int(total_twips * weight / weight_total)) for weight in weights]
    scale = total_twips / sum(widths)
    widths = [int(width * scale) for width in widths]
    widths[-1] += total_twips - sum(widths)
    for row_index, source_row in enumerate(padded):
        for column_index, text in enumerate(source_row):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _append_inline(paragraph, text)
            for run in paragraph.runs:
                _set_run_font(run, "Arial", 8, IVORY if row_index == 0 else BODY, row_index == 0)
            if row_index == 0:
                _set_cell_shading(cell, NAVY)
            elif row_index % 2 == 1:
                _set_cell_shading(cell, "F6F2E9")
    _set_repeat_table_header(table.rows[0])
    _set_table_geometry(table, widths)


def _remove_body_after(document: DocumentObject, paragraph_index: int) -> None:
    paragraphs = document.paragraphs
    if paragraph_index >= len(paragraphs):
        raise ValueError(f"DOCX cover break index {paragraph_index} exceeds paragraph count")
    for paragraph in paragraphs[paragraph_index + 1 :]:
        paragraph._element.getparent().remove(paragraph._element)
    for table in list(document.tables):
        table._element.getparent().remove(table._element)


def _append_blocks(document: DocumentObject, blocks: list[tuple[str, object]]) -> None:
    for kind, value in blocks:
        if kind in {"heading2", "heading3"}:
            paragraph = document.add_paragraph(style="Heading 1" if kind == "heading2" else "Heading 2")
            _append_inline(paragraph, str(value))
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _append_inline(paragraph, str(value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind in {"ordered", "unordered"}:
            style = "List Number" if kind == "ordered" else "List Bullet"
            for item in value:
                paragraph = document.add_paragraph(style=style)
                _append_inline(paragraph, str(item))
        elif kind == "quote":
            for item in value:
                paragraph = document.add_paragraph(style="Quote")
                _append_inline(paragraph, str(item))
        elif kind == "table":
            _add_table(document, value)
        else:
            raise ValueError(f"Unsupported DOCX rendering block: {kind}")


def _configure_document(document: DocumentObject, spec: DocxSpec) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    properties = document.core_properties
    properties.title = spec.pdf_spec.title
    properties.subject = "AI-IDP / AegisTrace"
    properties.author = "Pierre-Edward Procyk"
    properties.last_modified_by = "Cognitive Industries — Les Industries Cognitives"
    properties.category = spec.pdf_spec.classification
    properties.comments = "© 2026 Pierre-Edward Procyk. All rights reserved."
    properties.created = datetime(2026, 8, 1)
    properties.modified = datetime(2026, 8, 1)
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _plain_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return re.sub(r"[*`]", "", text)


def _document_text(document: DocumentObject) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def _validate_docx(path: Path, spec: DocxSpec, blocks: list[tuple[str, object]]) -> dict[str, object]:
    document = Document(path)
    text = _document_text(document)
    required = (*spec.pdf_spec.required_text, spec.pdf_spec.document_id)
    missing = [value for value in required if value not in text]
    stale = [
        value
        for value in spec.pdf_spec.forbidden_text
        if re.search(rf"(?<!\d){re.escape(value)}(?!\d)", text)
    ]
    uncovered: list[str] = []
    normalized = re.sub(r"\s+", " ", text).strip()
    for kind, value in blocks:
        values = value if kind in {"ordered", "unordered", "quote", "table"} else (value,)
        flattened = [cell for row in values for cell in row] if kind == "table" else values
        for item in flattened:
            expected = re.sub(r"\s+", " ", _plain_markdown(str(item))).strip()
            if expected and expected not in normalized:
                uncovered.append(expected[:160])
    if missing or stale or uncovered:
        raise ValueError(f"{path}: missing={missing}, stale={stale}, uncovered={uncovered[:5]}")
    if not any(section.start_type is not None for section in document.sections):
        raise ValueError(f"{path}: no valid section geometry")
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "characters": len(text),
    }


def build_docx(spec: DocxSpec) -> dict[str, object]:
    if not spec.reference.is_file():
        raise FileNotFoundError(spec.reference)
    if not BRAND_LOGO.is_file():
        raise FileNotFoundError(BRAND_LOGO)
    markdown = spec.pdf_spec.source.read_text(encoding="utf-8-sig")
    blocks = _parse_markdown(markdown, spec.pdf_spec.cover_heading_count)
    document = Document(spec.reference)
    if len(document.paragraphs) <= spec.cover_page_break_index:
        raise ValueError(f"{spec.reference}: reference cover structure changed")
    for index, value in spec.cover_updates:
        _replace_paragraph_text(document.paragraphs[index], value)
    _remove_body_after(document, spec.cover_page_break_index)
    _configure_document(document, spec)
    _append_blocks(document, blocks)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{spec.reference.stem}.", suffix=".tmp.docx", dir=spec.reference.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(temporary)
        validation = _validate_docx(temporary, spec, blocks)
        os.replace(temporary, spec.reference)
    finally:
        temporary.unlink(missing_ok=True)
    return {"key": spec.key, "output": str(spec.reference), **validation}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", choices=tuple(PDF_DOCUMENTS), action="append")
    parser.add_argument("--docx", choices=tuple(DOCX_DOCUMENTS), action="append")
    parser.add_argument("--all", action="store_true", help="Build every configured PDF and DOCX derivative")
    args = parser.parse_args()
    if not args.all and not args.pdf and not args.docx:
        parser.error("select --all, --pdf, or --docx")
    pdf_keys = list(PDF_DOCUMENTS) if args.all else (args.pdf or [])
    docx_keys = list(DOCX_DOCUMENTS) if args.all else (args.docx or [])
    body_font, bold_font = _register_fonts()
    for key in pdf_keys:
        result = build_document(PDF_DOCUMENTS[key], body_font, bold_font)
        print(f"pdf:{result['key']}: {result['pages']} pages, {result['characters']} characters -> {result['output']}")
    for key in docx_keys:
        result = build_docx(DOCX_DOCUMENTS[key])
        print(
            f"docx:{result['key']}: {result['paragraphs']} paragraphs, "
            f"{result['tables']} tables, {result['characters']} characters -> {result['output']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
